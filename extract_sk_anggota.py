#!/usr/bin/env python3
"""
extract_sk_anggota.py — Ekstrak tabel "Daftar Anggota" dari PDF SK (hasil SCAN)
memakai Claude vision (Anthropic API), output file Excel siap import.

Contoh:
    python extract_sk_anggota.py --input SK_KTH_SUMBER_JATI.pdf --pages 9-12
    python extract_sk_anggota.py --input SK.pdf --pages 9-12 --output anggota.xlsx --dpi 400

Kebutuhan: Python 3.10+, library `anthropic`, `PyMuPDF` (fitz), `openpyxl`.
    pip install anthropic pymupdf openpyxl

API key diambil dari environment variable ANTHROPIC_API_KEY (jangan hardcode).
"""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import sys
import time

DEFAULT_MODEL = "claude-sonnet-5"
DEFAULT_DPI = 300
DEFAULT_DELAY = 2.0
DEFAULT_MAX_TOKENS = 4096

PROMPT = """\
Anda membaca halaman lampiran SK (Surat Keputusan) Kemitraan Kehutanan / Perhutanan Sosial.
Halaman ini berisi tabel "DAFTAR ANGGOTA" dengan kolom: NO | NAMA | L/P | NIK | DESA | KECAMATAN.

Tugas: baca SELURUH baris tabel anggota pada gambar ini dan kembalikan HANYA sebuah JSON array
(tanpa teks pembuka/penutup, tanpa markdown code fence), dengan struktur tiap elemen:

{"no": 1, "nama": "SUKADI", "nik": "3522033112630018", "lp": "L", "desa": "BONDOL", "kecamatan": "NGAMBON"}

Aturan WAJIB:
1. Baca NIK dengan sangat teliti digit per digit — NIK harus selalu tepat 16 digit angka.
   Kalau ada digit yang tidak yakin / tidak jelas terbaca, isi field "nik" dengan string kosong ""
   dan tambahkan field "catatan": "NIK tidak terbaca jelas, perlu verifikasi manual" pada baris
   tersebut. JANGAN mengarang atau menebak digit yang tidak jelas.
2. Kalau ada baris tabel yang nomornya tidak berurutan atau ada baris tanpa data lengkap,
   tetap sertakan apa adanya, jangan dilewati.
3. Kalau halaman ini TIDAK memuat tabel daftar anggota sama sekali, kembalikan array kosong: []
4. Kembalikan HANYA JSON yang valid, tidak ada teks lain di luar JSON.
"""


# --------------------------------------------------------------------------- util

def parse_pages(spec: str, total: int) -> list[int]:
    """Ubah spesifikasi rentang 1-based ("9-12", "5", "9,10,12") jadi list nomor halaman.

    Melempar ValueError dengan pesan jelas bila format salah / di luar jangkauan.
    """
    if not spec or not spec.strip():
        return list(range(1, total + 1))
    pages: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            bounds = part.split("-")
            if len(bounds) != 2:
                raise ValueError(f"Rentang halaman tidak valid: '{part}'. Contoh benar: '9-12'.")
            try:
                start, end = int(bounds[0]), int(bounds[1])
            except ValueError:
                raise ValueError(f"Rentang halaman tidak valid: '{part}'. Gunakan angka, contoh: '9-12'.")
            if start < 1 or end < 1 or end < start:
                raise ValueError(f"Rentang halaman tidak valid: '{part}'. Contoh benar: '9-12'.")
            pages.extend(range(start, end + 1))
        else:
            try:
                pages.append(int(part))
            except ValueError:
                raise ValueError(f"Nomor halaman tidak valid: '{part}'. Gunakan angka / rentang, contoh: '9-12'.")
    # buang duplikat, urutkan
    pages = sorted(set(pages))
    bad = [p for p in pages if p < 1 or p > total]
    if bad:
        raise ValueError(
            f"Halaman {bad} di luar jangkauan: PDF hanya punya {total} halaman (1-{total})."
        )
    if not pages:
        raise ValueError("Rentang halaman kosong.")
    return pages


def render_page_png(doc, page_index_0based: int, dpi: int) -> bytes:
    """Render satu halaman PDF menjadi bytes PNG (di memori, tanpa file temp)."""
    page = doc[page_index_0based]
    pix = page.get_pixmap(dpi=dpi)
    return pix.tobytes("png")


def extract_page_text(client, model: str, img_png: bytes, max_tokens: int) -> str:
    """Kirim satu gambar halaman ke Claude vision, kembalikan teks respons mentah."""
    b64 = base64.b64encode(img_png).decode("ascii")
    resp = client.messages.create(
        model=model,
        max_tokens=max_tokens,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image",
                 "source": {"type": "base64", "media_type": "image/png", "data": b64}},
                {"type": "text", "text": PROMPT},
            ],
        }],
    )
    return "".join(getattr(block, "text", "") for block in resp.content)


def parse_json_response(text: str, page_no: int) -> list[dict] | None:
    """Parse respons model menjadi list dict. Kembalikan None bila bukan JSON valid
    (penelepon WAJIB lanjut ke halaman berikut, bukan crash)."""
    t = text.strip()
    # buang markdown code fence bila ada
    t = re.sub(r"^```(?:json)?\s*", "", t)
    t = re.sub(r"\s*```$", "", t)
    try:
        data = json.loads(t)
    except json.JSONDecodeError:
        # fallback: ambil blok array [...] pertama
        m = re.search(r"\[.*\]", t, re.DOTALL)
        if not m:
            print(f"  ! Halaman {page_no}: respons bukan JSON valid, dilewati.", flush=True)
            return None
        try:
            data = json.loads(m.group(0))
        except json.JSONDecodeError:
            print(f"  ! Halaman {page_no}: respons bukan JSON valid, dilewati.", flush=True)
            return None
    if not isinstance(data, list):
        print(f"  ! Halaman {page_no}: JSON bukan array, dilewati.", flush=True)
        return None
    rows = [r for r in data if isinstance(r, dict)]
    if len(rows) != len(data):
        print(f"  ! Halaman {page_no}: {len(data) - len(rows)} item non-objek diabaikan.", flush=True)
    return rows


def bersihkan_nik(nik) -> str:
    """Ambil digit saja dari NIK (buang spasi/titik/tanda baca hasil baca)."""
    return re.sub(r"\D", "", str(nik or ""))


def tambah_catatan(row: dict, pesan: str) -> None:
    lama = (row.get("catatan") or "").strip()
    row["catatan"] = f"{lama} | {pesan}" if lama else pesan


def normalisasi_baris(raw: dict) -> dict:
    return {
        "no": raw.get("no", ""),
        "nama": str(raw.get("nama", "") or "").strip().upper(),
        "nik": bersihkan_nik(raw.get("nik", "")),
        "lp": str(raw.get("lp", "") or "").strip().upper(),
        "desa": str(raw.get("desa", "") or "").strip().upper(),
        "kecamatan": str(raw.get("kecamatan", "") or "").strip().upper(),
        "catatan": str(raw.get("catatan", "") or "").strip(),
    }


def validasi_rows(rows: list[dict]) -> tuple[list[dict], int]:
    """Validasi NIK 16 digit + tandai duplikat. Kembalikan (rows, jumlah_butuh_manual)."""
    bersih = [normalisasi_baris(r) for r in rows]
    for r in bersih:
        if len(r["nik"]) != 16 or not r["nik"].isdigit():
            tambah_catatan(r, "NIK tidak valid (bukan 16 digit)")
    pertama: dict[str, object] = {}
    for r in bersih:
        nik = r["nik"]
        if not nik:
            continue
        if nik in pertama:
            tambah_catatan(r, f"NIK duplikat (sama dengan baris no {pertama[nik]})")
        else:
            pertama[nik] = r["no"]
    need = sum(1 for r in bersih if (r.get("catatan") or "").strip())
    return bersih, need


def tulis_excel(rows: list[dict], path: str) -> None:
    """Tulis hasil ke .xlsx; cell Catatan yang terisi diberi highlight kuning."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Anggota SK"
    header = ["No", "Nama", "NIK", "L/P", "Desa", "Kecamatan", "Catatan"]
    hdr_fill = PatternFill("solid", fgColor="D9EAD3")
    warn_fill = PatternFill("solid", fgColor="FFFF00")
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([r["no"], r["nama"], r["nik"], r["lp"], r["desa"], r["kecamatan"], r["catatan"]])
        n = ws.max_row
        ws.cell(n, 3).number_format = "@"  # NIK sebagai teks (anti notasi ilmiah)
        if (r["catatan"] or "").strip():
            ws.cell(n, 7).fill = warn_fill
    for col, w in zip("ABCDEFG", [6, 28, 20, 6, 20, 20, 55]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    wb.save(path)


def gagal(pesan: str, kode: int = 1) -> "NoReturn":
    print(f"ERROR: {pesan}", file=sys.stderr)
    raise SystemExit(kode)


# --------------------------------------------------------------------------- main

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Ekstrak tabel Daftar Anggota dari PDF SK hasil scan via Claude vision ke Excel.",
        epilog="Contoh: python extract_sk_anggota.py --input SK.pdf --pages 9-12 --output anggota.xlsx",
    )
    p.add_argument("--input", required=True, help="Path file PDF SK.")
    p.add_argument("--pages", default="",
                   help="Rentang halaman 1-based, cth '9-12' atau '9,10,12'. "
                        "Dikosongkan = seluruh halaman (halaman tanpa tabel dilewati).")
    p.add_argument("--output", default="",
                   help="Path file .xlsx output. Default: <nama_input>_anggota.xlsx.")
    p.add_argument("--model", default=DEFAULT_MODEL, help=f"Model Anthropic (default: {DEFAULT_MODEL}).")
    p.add_argument("--dpi", type=int, default=DEFAULT_DPI,
                   help=f"Resolusi render halaman, dpi (default: {DEFAULT_DPI}).")
    p.add_argument("--delay", type=float, default=DEFAULT_DELAY,
                   help=f"Jeda detik antar panggilan API (default: {DEFAULT_DELAY}).")
    p.add_argument("--max-tokens", type=int, default=DEFAULT_MAX_TOKENS,
                   help=f"Batas token output per halaman (default: {DEFAULT_MAX_TOKENS}).")
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)

    if not os.path.isfile(args.input):
        gagal(f"File PDF tidak ditemukan: {args.input}")
    api_key = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if not api_key:
        gagal("ANTHROPIC_API_KEY belum diset. Contoh (PowerShell): "
              '$env:ANTHROPIC_API_KEY="sk-ant-..."')
    if args.dpi < 72 or args.dpi > 600:
        gagal(f"--dpi di luar wajar (72-600): {args.dpi}")

    try:
        try:
            import pymupdf as fitz  # nama modern (tanpa warning deprecated)
        except ImportError:
            import fitz  # PyMuPDF lama
    except ImportError:
        gagal("Library PyMuPDF belum terpasang. Jalankan: pip install pymupdf")
    try:
        import anthropic
    except ImportError:
        gagal("Library anthropic belum terpasang. Jalankan: pip install anthropic")

    try:
        doc = fitz.open(args.input)
    except Exception as e:
        gagal(f"Tidak bisa membuka PDF '{args.input}': {e}")
    total = doc.page_count
    if total == 0:
        gagal(f"PDF '{args.input}' tidak berisi halaman.")

    try:
        pages = parse_pages(args.pages, total)
    except ValueError as e:
        gagal(str(e))

    output = args.output.strip() or (
        os.path.splitext(args.input)[0] + "_anggota.xlsx"
    )

    client = anthropic.Anthropic()  # baca ANTHROPIC_API_KEY dari environment
    semua: list[dict] = []
    gagal_halaman = 0
    hlm_str = f"{pages[0]}-{pages[-1]}" if len(pages) > 1 else f"halaman {pages[0]}"
    print(f"PDF: {args.input} ({total} halaman) -> memproses {len(pages)} halaman: {hlm_str}")
    for i, pno in enumerate(pages):
        print(f"[{i + 1}/{len(pages)}] Halaman {pno} dirender {args.dpi}dpi ...", flush=True)
        try:
            img = render_page_png(doc, pno - 1, args.dpi)
        except Exception as e:
            print(f"  ! Halaman {pno}: gagal render ({e}), dilewati.", flush=True)
            gagal_halaman += 1
            continue
        try:
            teks = extract_page_text(client, args.model, img, args.max_tokens)
        except Exception as e:
            print(f"  ! Halaman {pno}: panggilan API gagal ({e}), dilewati.", flush=True)
            gagal_halaman += 1
            continue
        rows = parse_json_response(teks, pno)
        if rows is None:
            gagal_halaman += 1
            continue
        print(f"  -> {len(rows)} baris.", flush=True)
        semua.extend(rows)
        if i < len(pages) - 1 and args.delay > 0:
            time.sleep(args.delay)
    doc.close()

    if not semua and gagal_halaman:
        gagal("Tidak ada baris berhasil diekstrak (semua halaman gagal). Periksa koneksi/API key.")
    final, need = validasi_rows(semua)
    try:
        tulis_excel(final, output)
    except Exception as e:
        gagal(f"Gagal menulis Excel '{output}': {e}")

    print("-" * 50)
    print(f"Total baris diekstrak : {len(final)}")
    print(f"Butuh verifikasi manual: {need} baris (kolom Catatan, highlight kuning)")
    if gagal_halaman:
        print(f"Halaman gagal diproses: {gagal_halaman} (sudah dilewati, tidak menghentikan skrip)")
    print(f"File output           : {output}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as e:  # pengaman terakhir: pesan bersih, tanpa traceback
        print(f"ERROR tak terduga: {e}", file=sys.stderr)
        raise SystemExit(1)
