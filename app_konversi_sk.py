import ctypes
import os
import sys

def _setup_windows_dlls():
    if sys.platform != "win32":
        return
    dirs = []
    if getattr(sys, "frozen", False):
        base_dir = os.path.dirname(sys.executable)
        dirs = [
            base_dir,
            os.path.join(base_dir, "_internal"),
            os.path.join(base_dir, "_internal", "onnxruntime", "capi"),
            os.path.join(base_dir, "onnxruntime", "capi"),
            os.path.join(base_dir, "_internal", "cv2"),
            os.path.join(base_dir, "cv2"),
        ]
    else:
        try:
            import onnxruntime
            dirs.append(os.path.join(os.path.dirname(onnxruntime.__file__), "capi"))
        except Exception:
            pass

    for d in dirs:
        if os.path.isdir(d):
            try:
                os.add_dll_directory(d)
            except Exception:
                pass
            os.environ["PATH"] = d + ";" + os.environ.get("PATH", "")

    # Preload essential VC++ and ONNX Runtime DLLs
    preload_names = [
        "vcomp140.dll",
        "msvcp140.dll",
        "msvcp140_1.dll",
        "msvcp140_2.dll",
        "vcruntime140.dll",
        "vcruntime140_1.dll",
        "onnxruntime.dll",
        "onnxruntime_providers_shared.dll"
    ]
    for d in dirs:
        for dll_name in preload_names:
            dll_path = os.path.join(d, dll_name)
            if os.path.isfile(dll_path):
                try:
                    ctypes.CDLL(dll_path)
                except Exception:
                    pass

_setup_windows_dlls()

import base64
import json
import re
import subprocess
import threading
import time
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import numpy as np
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import pymupdf


def cari_poppler() -> str | None:
    return None


def parse_page_ranges(input_str: str, total: int) -> list[int]:
    if not input_str.strip():
        return list(range(1, total + 1))
    pages = []
    for part in input_str.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            s, e = part.split("-", 1)
            pages.extend(range(int(s.strip()), int(e.strip()) + 1))
        else:
            pages.append(int(part))
    pages = sorted(set(pages))
    bad = [p for p in pages if p < 1 or p > total]
    if bad:
        raise ValueError(f"Halaman {bad} di luar jangkauan (1-{total}).")
    return pages


def deteksi_lp_dari_nik(nik: str) -> str:
    nik = str(nik or "").strip()
    if len(nik) == 16 and nik.isdigit():
        tgl = int(nik[6:8])
        if 1 <= tgl <= 31:
            return "L"
        elif 41 <= tgl <= 71:
            return "P"
    return ""


def extract_offline_core(pdf_path: str, pages: list[int], log_fn, progress_fn) -> list[dict]:
    try:
        from rapidocr_onnxruntime import RapidOCR
        ocr = RapidOCR()
    except Exception as e:
        raise RuntimeError(
            f"Gagal memuat engine OCR lokal ({e}).\n"
            "Solusi: Unduh dan instal 'Microsoft Visual C++ 2015-2022 Redistributable (x64)' "
            "dari link resmi Microsoft:\nhttps://aka.ms/vs/17/release/vc_redist.x64.exe"
        )
    doc = pymupdf.open(pdf_path)
    semua_baris = []
    total_p = len(pages)

    for idx, pno in enumerate(pages, 1):
        log_fn(f"[{idx}/{total_p}] Merender & OCR Halaman {pno} (300 dpi)...")
        progress_fn(idx / total_p * 0.9)

        pix = doc[pno - 1].get_pixmap(dpi=300)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        res, _ = ocr(img)
        res = res or []

        header_y = 0
        for r in res:
            text = r[1].upper()
            box = np.array(r[0])
            yc = box[:, 1].mean()
            if any(k in text for k in ["NAMA", "NIK", "KECAMATAN", "DESA", "L/P"]):
                if yc > header_y:
                    header_y = yc
        if header_y == 0:
            header_y = 300

        table_boxes = []
        for r in res:
            box = np.array(r[0])
            yc = box[:, 1].mean()
            xc = box[:, 0].mean()
            if yc > header_y + 30:
                table_boxes.append({
                    "text": r[1].strip(),
                    "score": float(r[2]),
                    "xc": xc,
                    "yc": yc,
                    "box": box
                })

        nik_boxes = []
        for b in table_boxes:
            clean_digits = re.sub(r"\D", "", b["text"])
            if 1100 <= b["xc"] <= 1550 and len(clean_digits) in (15, 16, 17):
                nik_boxes.append((b["yc"], clean_digits, b))

        nik_boxes.sort(key=lambda x: x[0])

        for yc_anchor, nik_digits, _ in nik_boxes:
            row_items = [b for b in table_boxes if abs(b["yc"] - yc_anchor) <= 35]

            no_val = ""
            nama_parts = []
            desa_parts = []
            kec_parts = []
            lp_ocr = ""

            for item in sorted(row_items, key=lambda x: x["xc"]):
                xc = item["xc"]
                txt = item["text"]
                if xc < 250:
                    if re.match(r"^\d+$", txt):
                        no_val = txt
                elif 250 <= xc < 850:
                    if not txt.isdigit():
                        nama_parts.append(txt)
                elif 850 <= xc < 1100:
                    c_lp = txt.upper().strip()
                    if c_lp in ("L", "P"):
                        lp_ocr = c_lp
                elif 1100 <= xc < 1550:
                    pass
                elif 1550 <= xc < 1950:
                    desa_parts.append(txt)
                elif xc >= 1950:
                    kec_parts.append(txt)

            nama_val = " ".join(nama_parts).strip().upper()
            desa_val = " ".join(desa_parts).strip().upper()
            kec_val = " ".join(kec_parts).strip().upper()
            lp_final = lp_ocr or deteksi_lp_dari_nik(nik_digits)

            semua_baris.append({
                "no": no_val,
                "nama": nama_val,
                "nik": nik_digits,
                "lp": lp_final,
                "desa": desa_val,
                "kecamatan": kec_val,
                "page": pno
            })

    doc.close()
    return semua_baris


def extract_online_core(pdf_path: str, pages: list[int], api_key: str, log_fn, progress_fn) -> list[dict]:
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)
    doc = pymupdf.open(pdf_path)
    semua_baris = []
    total_p = len(pages)

    prompt = """Anda membaca halaman lampiran SK Kemitraan Kehutanan/Perhutanan Sosial yang memuat tabel "DAFTAR ANGGOTA".
Tugas: baca SELURUH baris tabel dan kembalikan HANYA JSON array (tanpa markdown fence, tanpa teks lain):
[
  {"no": 1, "nama": "SUKADI", "nik": "3522033112630018", "lp": "L", "desa": "BONDOL", "kecamatan": "NGAMBON"}
]
Aturan penting:
1. Baca NIK sangat teliti 16 digit. Kalau ada digit yang buram/ragu, isi "" dan beri catatan: "NIK tidak terbaca jelas, perlu verifikasi manual".
2. Kalau halaman tidak ada tabel anggota, kembalikan [].
"""

    for idx, pno in enumerate(pages, 1):
        log_fn(f"[{idx}/{total_p}] Mengirim gambar Halaman {pno} ke Claude Vision...")
        progress_fn(idx / total_p * 0.9)

        pix = doc[pno - 1].get_pixmap(dpi=300)
        img_png = pix.tobytes("png")
        b64 = base64.b64encode(img_png).decode("ascii")

        try:
            resp = client.messages.create(
                model="claude-sonnet-5",
                max_tokens=4096,
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": b64}},
                        {"type": "text", "text": prompt}
                    ]
                }]
            )
            raw_text = "".join(getattr(b, "text", "") for b in resp.content).strip()
            raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text)
            raw_text = re.sub(r"\s*```$", "", raw_text)
            try:
                data = json.loads(raw_text)
            except Exception:
                m = re.search(r"\[.*\]", raw_text, re.DOTALL)
                data = json.loads(m.group(0)) if m else []

            if isinstance(data, list):
                for item in data:
                    if isinstance(item, dict):
                        semua_baris.append({
                            "no": item.get("no", ""),
                            "nama": str(item.get("nama", "")).strip().upper(),
                            "nik": re.sub(r"\D", "", str(item.get("nik", ""))),
                            "lp": str(item.get("lp", "")).strip().upper(),
                            "desa": str(item.get("desa", "")).strip().upper(),
                            "kecamatan": str(item.get("kecamatan", "")).strip().upper(),
                            "catatan": str(item.get("catatan", "")).strip()
                        })
                log_fn(f"  -> Halaman {pno}: didapat {len(data)} baris.")
        except Exception as e:
            log_fn(f"  ! Halaman {pno} gagal: {e}")

        time.sleep(1.0)

    doc.close()
    return semua_baris


def validasi_dan_simpan_excel(rows: list[dict], output_path: str) -> tuple[int, int]:
    for idx, r in enumerate(rows, start=1):
        if not r.get("no") or not str(r["no"]).isdigit():
            r["no"] = idx
        catatan_list = []
        lama = r.get("catatan", "").strip()
        if lama:
            catatan_list.append(lama)
        if len(r["nik"]) != 16:
            catatan_list.append(f"NIK tidak 16 digit ({len(r['nik'])} digit)")
        if not r["nama"]:
            catatan_list.append("Nama kosong")
        r["catatan"] = " | ".join(catatan_list)

    seen = {}
    for r in rows:
        nik = r["nik"]
        if nik and len(nik) == 16:
            if nik in seen:
                pesan = f"NIK duplikat (sama no {seen[nik]})"
                r["catatan"] = f"{r['catatan']} | {pesan}" if r["catatan"] else pesan
            else:
                seen[nik] = r["no"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Daftar Anggota"

    header = ["No", "Nama", "NIK", "L/P", "Desa", "Kecamatan", "Catatan"]
    ws.append(header)

    hdr_fill = PatternFill("solid", fgColor="D9EAD3")
    warn_fill = PatternFill("solid", fgColor="FFFF00")

    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([r["no"], r["nama"], r["nik"], r["lp"], r["desa"], r["kecamatan"], r["catatan"]])
        n = ws.max_row
        ws.cell(n, 3).number_format = "@"
        if r["catatan"]:
            ws.cell(n, 7).fill = warn_fill

    for col, w in zip("ABCDEFG", [6, 28, 20, 6, 20, 20, 45]):
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    wb.save(output_path)

    need_manual = sum(1 for r in rows if r["catatan"])
    return len(rows), need_manual


# --------------------------------------------------------------------------- GUI Interface

class AppKonversiSK(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Konversi PDF SK ke Excel Anggota — Pupuk PS")
        self.geometry("760x650")
        self.minsize(680, 560)

        # Style & Font
        self.configure(bg="#F4F6F9")
        self.style = ttk.Style(self)
        self.style.theme_use("clam")

        self.last_output_file = ""
        self.is_running = False

        self._build_ui()

    def _build_ui(self):
        # Header Banner
        header = tk.Frame(self, bg="#1E3A8A", height=70)
        header.pack(fill="x")
        header.pack_propagate(False)

        title = tk.Label(header, text="📄 Konverter PDF SK ke Data Excel Anggota",
                         font=("Segoe UI", 15, "bold"), fg="#FFFFFF", bg="#1E3A8A")
        title.pack(anchor="w", padx=20, pady=(12, 2))

        subtitle = tk.Label(header, text="Ekstrak tabel Daftar Anggota dari PDF SK hasil scan secara cepat & akurat.",
                            font=("Segoe UI", 9), fg="#93C5FD", bg="#1E3A8A")
        subtitle.pack(anchor="w", padx=20)

        # Main Container
        main_frame = tk.Frame(self, bg="#F4F6F9", padx=20, pady=15)
        main_frame.pack(fill="both", expand=True)

        # 1. Card Input File
        card_file = tk.LabelFrame(main_frame, text=" 1. Pilih File Dokumen SK ",
                                  font=("Segoe UI", 10, "bold"), bg="#FFFFFF", padx=15, pady=12)
        card_file.pack(fill="x", pady=(0, 10))

        row1 = tk.Frame(card_file, bg="#FFFFFF")
        row1.pack(fill="x", pady=4)

        tk.Label(row1, text="File PDF:", font=("Segoe UI", 9, "bold"), width=12, anchor="w", bg="#FFFFFF").pack(side="left")
        self.ent_pdf = ttk.Entry(row1)
        self.ent_pdf.pack(side="left", fill="x", expand=True, padx=(0, 8))
        btn_browse_pdf = ttk.Button(row1, text="Pilih PDF...", command=self._browse_pdf)
        btn_browse_pdf.pack(side="right")

        self.lbl_pdf_info = tk.Label(card_file, text="Belum ada file dipilih", font=("Segoe UI", 8, "italic"), fg="#6B7280", bg="#FFFFFF")
        self.lbl_pdf_info.pack(anchor="w", padx=95)

        # 2. Card Pengaturan Ekstraksi
        card_opt = tk.LabelFrame(main_frame, text=" 2. Pengaturan Ekstraksi ",
                                 font=("Segoe UI", 10, "bold"), bg="#FFFFFF", padx=15, pady=12)
        card_opt.pack(fill="x", pady=(0, 10))

        # Rentang Halaman
        row_pg = tk.Frame(card_opt, bg="#FFFFFF")
        row_pg.pack(fill="x", pady=4)
        tk.Label(row_pg, text="Halaman:", font=("Segoe UI", 9, "bold"), width=12, anchor="w", bg="#FFFFFF").pack(side="left")
        self.ent_pages = ttk.Entry(row_pg, width=20)
        self.ent_pages.insert(0, "9-12")
        self.ent_pages.pack(side="left", padx=(0, 10))
        tk.Label(row_pg, text="Contoh: 9-12 (kosongkan jika ingin scan semua halaman)",
                 font=("Segoe UI", 8), fg="#6B7280", bg="#FFFFFF").pack(side="left")

        # Pilihan Mode
        row_mode = tk.Frame(card_opt, bg="#FFFFFF")
        row_mode.pack(fill="x", pady=(10, 4))
        tk.Label(row_mode, text="Mode Engine:", font=("Segoe UI", 9, "bold"), width=12, anchor="w", bg="#FFFFFF").pack(side="left")

        self.mode_var = tk.StringVar(value="offline")
        rb_off = ttk.Radiobutton(row_mode, text="100% Offline (RapidOCR - Cepat & Tanpa Internet)",
                                 variable=self.mode_var, value="offline", command=self._toggle_mode)
        rb_off.pack(side="left", padx=(0, 15))

        rb_on = ttk.Radiobutton(row_mode, text="Online AI (Claude Vision - Akurasi NIK Maksimal)",
                                variable=self.mode_var, value="online", command=self._toggle_mode)
        rb_on.pack(side="left")

        # API Key frame (sembunyi jika offline)
        self.frame_api = tk.Frame(card_opt, bg="#FFFFFF")
        self.frame_api.pack(fill="x", pady=(6, 0))
        tk.Label(self.frame_api, text="API Key:", font=("Segoe UI", 9, "bold"), width=12, anchor="w", bg="#FFFFFF").pack(side="left")
        self.ent_api = ttk.Entry(self.frame_api, show="•")
        # Load default from env if exists
        default_key = os.environ.get("ANTHROPIC_API_KEY", "")
        if default_key:
            self.ent_api.insert(0, default_key)
        self.ent_api.pack(side="left", fill="x", expand=True)
        self.frame_api.pack_forget()

        # Output Excel
        row_out = tk.Frame(card_opt, bg="#FFFFFF")
        row_out.pack(fill="x", pady=(10, 4))
        tk.Label(row_out, text="Output Excel:", font=("Segoe UI", 9, "bold"), width=12, anchor="w", bg="#FFFFFF").pack(side="left")
        self.ent_output = ttk.Entry(row_out)
        self.ent_output.pack(side="left", fill="x", expand=True, padx=(0, 8))
        btn_browse_out = ttk.Button(row_out, text="Pilih Lokasi...", command=self._browse_output)
        btn_browse_out.pack(side="right")

        # 3. Action & Progress
        act_frame = tk.Frame(main_frame, bg="#F4F6F9")
        act_frame.pack(fill="x", pady=6)

        self.btn_run = tk.Button(act_frame, text="▶  MULAI KONVERSI KE EXCEL",
                                 font=("Segoe UI", 11, "bold"), bg="#10B981", fg="#FFFFFF",
                                 activebackground="#059669", activeforeground="#FFFFFF",
                                 relief="flat", padx=20, pady=8, cursor="hand2", command=self._start_conversion)
        self.btn_run.pack(side="left")

        self.btn_open_excel = tk.Button(act_frame, text="📊 Buka Hasil Excel",
                                        font=("Segoe UI", 10), bg="#E5E7EB", fg="#1F2937",
                                        relief="flat", padx=15, pady=8, state="disabled",
                                        cursor="hand2", command=self._open_excel)
        self.btn_open_excel.pack(side="left", padx=10)

        self.progress = ttk.Progressbar(act_frame, orient="horizontal", mode="determinate")
        self.progress.pack(side="right", fill="x", expand=True, padx=(10, 0))

        # 4. Log Console Box
        card_log = tk.LabelFrame(main_frame, text=" Status & Log Proses ",
                                 font=("Segoe UI", 9, "bold"), bg="#FFFFFF", padx=8, pady=8)
        card_log.pack(fill="both", expand=True, pady=(10, 0))

        self.txt_log = tk.Text(card_log, font=("Consolas", 9), bg="#1E293B", fg="#F8FAFC",
                               wrap="word", height=8, relief="flat")
        scroll = ttk.Scrollbar(card_log, orient="vertical", command=self.txt_log.yview)
        self.txt_log.configure(yscrollcommand=scroll.set)

        scroll.pack(side="right", fill="y")
        self.txt_log.pack(side="left", fill="both", expand=True)

        self._log("Siap. Pilih file PDF dan klik 'Mulai Konversi ke Excel'.")

    def _toggle_mode(self):
        if self.mode_var.get() == "online":
            self.frame_api.pack(fill="x", pady=(6, 0))
        else:
            self.frame_api.pack_forget()

    def _browse_pdf(self):
        f = filedialog.askopenfilename(
            title="Pilih File PDF SK",
            filetypes=[("PDF Document", "*.pdf"), ("All Files", "*.*")]
        )
        if f:
            self.ent_pdf.delete(0, "end")
            self.ent_pdf.insert(0, f)
            try:
                doc = pymupdf.open(f)
                count = doc.page_count
                doc.close()
                self.lbl_pdf_info.config(text=f"Total: {count} halaman | Siap diproses", fg="#059669")
                if not self.ent_output.get().strip():
                    default_out = os.path.splitext(f)[0] + "_anggota.xlsx"
                    self.ent_output.delete(0, "end")
                    self.ent_output.insert(0, default_out)
            except Exception as e:
                self.lbl_pdf_info.config(text=f"Gagal membaca PDF: {e}", fg="#DC2626")

    def _browse_output(self):
        f = filedialog.asksaveasfilename(
            title="Simpan File Excel Hasil",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")]
        )
        if f:
            self.ent_output.delete(0, "end")
            self.ent_output.insert(0, f)

    def _log(self, text: str):
        self.txt_log.insert("end", text + "\n")
        self.txt_log.see("end")

    def _set_progress(self, val: float):
        self.progress["value"] = int(val * 100)

    def _start_conversion(self):
        if self.is_running:
            return

        pdf_path = self.ent_pdf.get().strip()
        if not pdf_path or not os.path.isfile(pdf_path):
            messagebox.showerror("Error", "Silakan pilih file PDF yang valid terlebih dahulu.")
            return

        out_path = self.ent_output.get().strip()
        if not out_path:
            out_path = os.path.splitext(pdf_path)[0] + "_anggota.xlsx"
            self.ent_output.delete(0, "end")
            self.ent_output.insert(0, out_path)

        page_spec = self.ent_pages.get().strip()
        mode = self.mode_var.get()
        api_key = self.ent_api.get().strip() if mode == "online" else ""

        if mode == "online" and not api_key:
            messagebox.showerror("Error", "Mode Online AI membutuhkan Anthropic API Key.")
            return

        self.is_running = True
        self.btn_run.config(state="disabled", bg="#9CA3AF")
        self.btn_open_excel.config(state="disabled")
        self.progress["value"] = 0
        self.txt_log.delete("1.0", "end")

        thread = threading.Thread(
            target=self._worker_thread,
            args=(pdf_path, out_path, page_spec, mode, api_key),
            daemon=True
        )
        thread.start()

    def _worker_thread(self, pdf_path: str, out_path: str, page_spec: str, mode: str, api_key: str):
        try:
            doc = pymupdf.open(pdf_path)
            total_pages = doc.page_count
            doc.close()

            pages = parse_page_ranges(page_spec, total_pages)
            self._log(f"PDF: {os.path.basename(pdf_path)} ({total_pages} hal) -> Proses {len(pages)} hal ({pages[0]}..{pages[-1] if len(pages)>1 else pages[0]})")
            self._log(f"Engine: {'100% Offline (RapidOCR)' if mode == 'offline' else 'Online AI (Claude Vision)'}")

            start_t = time.time()
            if mode == "offline":
                rows = extract_offline_core(pdf_path, pages, self._log, self._set_progress)
            else:
                rows = extract_online_core(pdf_path, pages, api_key, self._log, self._set_progress)

            if not rows:
                raise RuntimeError("Tidak ada baris data anggota yang berhasil diekstrak.")

            self._log("Menyimpan & memvalidasi ke file Excel...")
            total_row, need_manual = validasi_dan_simpan_excel(rows, out_path)

            elapsed = time.time() - start_t
            self._set_progress(1.0)
            self.last_output_file = out_path

            self._log("-" * 50)
            self._log(f"SELESAI! Waktu: {elapsed:.1f} detik")
            self._log(f"Total baris diekstrak : {total_row}")
            self._log(f"Butuh verifikasi manual: {need_manual} baris (highlight kuning)")
            self._log(f"File tersimpan di     : {out_path}")

            self.after(0, lambda: self._on_success(total_row, need_manual, out_path))

        except Exception as e:
            self._log(f"\nERROR: {e}")
            self.after(0, lambda: messagebox.showerror("Konversi Gagal", f"Terjadi kesalahan saat memproses:\n{e}"))
        finally:
            self.after(0, self._on_finish)

    def _on_success(self, total: int, need: int, out_path: str):
        self.btn_open_excel.config(state="normal")
        msg = f"Berhasil mengekstrak {total} baris data anggota!\n"
        if need > 0:
            msg += f"\nCatatan: Ada {need} baris yang butuh verifikasi manual (diberi tanda highlight kuning pada file Excel)."
        msg += f"\n\nFile Output:\n{out_path}"
        messagebox.showinfo("Konversi Berhasil", msg)

    def _on_finish(self):
        self.is_running = False
        self.btn_run.config(state="normal", bg="#10B981")

    def _open_excel(self):
        if self.last_output_file and os.path.isfile(self.last_output_file):
            try:
                os.startfile(self.last_output_file)
            except Exception as e:
                messagebox.showerror("Error", f"Gagal membuka file Excel: {e}")


if __name__ == "__main__":
    app = AppKonversiSK()
    app.mainloop()
